from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'output' / 'Plano_de_Testes_Plataforma_PROPEG.docx'
LOGO = ROOT / 'static' / 'imagens' / 'logo_ufac.png'
PRODUCTION_URL = 'https://propeg-plataforma.onrender.com'
LOGIN_URL = f'{PRODUCTION_URL}/projetos/'

AZUL = '0B3D6E'
AZUL_CLARO = 'EAF2F8'
CINZA = 'F4F6F7'
BORDA = 'D9D9D9'


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for margin, value in (('top', top), ('start', start), ('bottom', bottom), ('end', end)):
        node = tc_mar.find(qn(f'w:{margin}'))
        if node is None:
            node = OxmlElement(f'w:{margin}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement('w:cantSplit')
    cant_split.set(qn('w:val'), 'true')
    tr_pr.append(cant_split)


def set_table_borders(table, color=BORDA, size='6'):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in('w:tblBorders')
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tbl_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = f'w:{edge}'
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), size)
        element.set(qn('w:color'), color)


def set_col_widths(table, widths):
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width


def format_table(table, widths=None, header=True):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    if widths:
        set_col_widths(table, widths)
    for row_index, row in enumerate(table.rows):
        set_row_cant_split(row)
        if header and row_index == 0:
            set_repeat_table_header(row)
        for cell in row.cells:
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if header and row_index == 0:
                set_cell_shading(cell, AZUL)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        run.font.bold = True
            elif row_index % 2 == 0:
                set_cell_shading(cell, AZUL_CLARO)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    for idx, header in enumerate(headers):
        table.rows[0].cells[idx].text = header
    for row_data in rows:
        row = table.add_row()
        for idx, value in enumerate(row_data):
            row.cells[idx].text = str(value)
    format_table(table, widths)
    doc.add_paragraph()
    return table


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run('Página ')
    run.font.size = Pt(9)
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instruction = OxmlElement('w:instrText')
    instruction.set(qn('xml:space'), 'preserve')
    instruction.text = 'PAGE'
    separate = OxmlElement('w:fldChar')
    separate.set(qn('w:fldCharType'), 'separate')
    result = OxmlElement('w:t')
    result.text = '1'
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    run._r.extend([begin, instruction, separate, result, end])


def add_bullet(doc, text, level=0):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.28 + (0.22 * level))
    paragraph.paragraph_format.first_line_indent = Inches(-0.22)
    paragraph.add_run(f'•  {text}')
    return paragraph


def add_numbered(doc, number, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.28)
    paragraph.paragraph_format.first_line_indent = Inches(-0.22)
    paragraph.add_run(f'{number}.  {text}')
    return paragraph


def add_label_paragraph(doc, label, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(f'{label}: ')
    run.bold = True
    paragraph.add_run(text)
    return paragraph


def add_case(doc, case):
    case_paragraphs = []
    heading = doc.add_paragraph(style='Heading 3')
    if case.get('page_break_before'):
        heading.paragraph_format.page_break_before = True
    case_paragraphs.append(heading)
    heading.add_run(f"{case['id']}  {case['title']}")
    profile = add_label_paragraph(doc, 'Perfil', case['profile'])
    priority = add_label_paragraph(doc, 'Prioridade', case.get('priority', 'Alta'))
    preconditions = add_label_paragraph(doc, 'Pré-condições', case['pre'])
    case_paragraphs.extend([profile, priority, preconditions])
    if case.get('data'):
        data = add_label_paragraph(doc, 'Dados de teste', case['data'])
        case_paragraphs.append(data)
    steps_label = doc.add_paragraph()
    case_paragraphs.append(steps_label)
    steps_label.paragraph_format.space_after = Pt(2)
    steps_label.add_run('Passos').bold = True
    for number, step in enumerate(case['steps'], start=1):
        case_paragraphs.append(add_numbered(doc, number, step))
    result_label = doc.add_paragraph()
    case_paragraphs.append(result_label)
    result_label.paragraph_format.space_after = Pt(2)
    result_label.add_run('Resultado esperado').bold = True
    for result in case['expected']:
        case_paragraphs.append(add_bullet(doc, result))

    # Mantém o cabeçalho e os metadados junto ao início dos passos, mas permite
    # que casos longos continuem na página seguinte sem criar grandes vazios.
    cabecalho = [heading, profile, priority, preconditions]
    if case.get('data'):
        cabecalho.append(data)
    cabecalho.append(steps_label)
    for paragraph in cabecalho:
        paragraph.paragraph_format.keep_with_next = True
    result_label.paragraph_format.keep_with_next = True
    for paragraph in case_paragraphs:
        paragraph.paragraph_format.keep_together = True
    doc.add_paragraph()


def add_cases_section(doc, title, intro, cases, page_break_before=False):
    heading = doc.add_heading(title, level=1)
    if page_break_before:
        heading.paragraph_format.page_break_before = True
    intro_paragraph = doc.add_paragraph(intro)
    heading.paragraph_format.keep_with_next = True
    intro_paragraph.paragraph_format.keep_with_next = True
    for case in cases:
        add_case(doc, case)


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Arial'
    normal._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
    normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    title_style = styles['Title']
    title_style.font.name = 'Arial'
    title_style._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
    title_style._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
    title_style.font.size = Pt(26)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_p_pr = title_style._element.get_or_add_pPr()
    title_borders = title_p_pr.find(qn('w:pBdr'))
    if title_borders is not None:
        title_p_pr.remove(title_borders)

    for style_name, size in (('Heading 1', 17), ('Heading 2', 14), ('Heading 3', 11.5)):
        style = styles[style_name]
        style.font.name = 'Arial'
        style._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
        style._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)

    add_page_number(section.footer.paragraphs[0])

    # Capa
    if LOGO.exists():
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(LOGO), width=Inches(1.25))
    doc.add_paragraph()
    title = doc.add_paragraph(style='Title')
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run('Plano de Testes da Plataforma de Projetos Institucionais')
    direct_title_borders = title._p.get_or_add_pPr().find(qn('w:pBdr'))
    if direct_title_borders is not None:
        title._p.get_or_add_pPr().remove(direct_title_borders)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(8)
    subtitle.paragraph_format.space_after = Pt(24)
    run = subtitle.add_run('Roteiro de validação funcional no ambiente de produção')
    run.font.size = Pt(15)
    run.font.bold = True

    metadata = doc.add_table(rows=4, cols=2)
    metadata.style = 'Table Grid'
    metadata.rows[0].cells[0].text = 'Sistema'
    metadata.rows[0].cells[1].text = 'Plataforma de Projetos Institucionais PROPEG UFAC'
    metadata.rows[1].cells[0].text = 'Versão do roteiro'
    metadata.rows[1].cells[1].text = '2.1'
    metadata.rows[2].cells[0].text = 'Data de referência'
    metadata.rows[2].cells[1].text = '16 de setembro de 2026'
    metadata.rows[3].cells[0].text = 'Ambiente'
    metadata.rows[3].cells[1].text = 'Produção no Render com dados temporários de amostra'
    format_table(metadata, [Inches(1.55), Inches(5.25)], header=False)
    for row in metadata.rows:
        set_cell_shading(row.cells[0], AZUL_CLARO)
        row.cells[0].paragraphs[0].runs[0].font.bold = True

    doc.add_paragraph()
    opening = doc.add_paragraph()
    opening.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    opening.add_run(
        'Este documento orienta a validação manual da plataforma publicada no Render, desde o cadastro '
        'e o login até a criação, tramitação, execução, finalização e consulta de projetos. '
        'Os casos utilizam dados temporários de amostra e cobrem os perfis Coordenador, Aluno, '
        'Gestor e o acesso restrito do Centro de Estudos. As operações realizadas alteram o banco '
        'de produção e devem permanecer limitadas às contas e aos registros descritos neste roteiro.'
    )
    doc.add_page_break()

    doc.add_heading('Como usar este roteiro', level=1)
    doc.add_paragraph(
        'Execute os casos na ordem sugerida quando quiser validar um fluxo completo. Casos isolados '
        'podem ser repetidos após recriar a base. Registre evidências em captura de tela e anote o '
        'resultado como Aprovado, Reprovado ou Bloqueado.'
    )
    add_bullet(doc, f'Acesse {LOGIN_URL}.')
    add_bullet(doc, 'Use somente as contas e os projetos de amostra identificados neste documento.')
    add_bullet(doc, 'Antes de cada fluxo completo, escolha um projeto de amostra que ainda não tenha sido alterado ou crie um novo projeto com uma das contas de coordenador.')
    add_bullet(doc, 'Para testar envio de e-mail, substitua temporariamente o e-mail de uma conta de amostra por uma caixa postal que você controle.')
    add_bullet(doc, 'Use um navegador atualizado e, para responsividade, teste também uma largura próxima de 390 px.')
    add_bullet(doc, 'As credenciais abaixo são temporárias e devem ser removidas ao fim da validação.')

    doc.add_heading('Critérios gerais de aprovação', level=2)
    criteria = [
        ('Aprovado', 'O comportamento observado corresponde integralmente ao resultado esperado.'),
        ('Reprovado', 'Há divergência funcional, mensagem incorreta, falha visual ou permissão indevida.'),
        ('Bloqueado', 'O teste não pôde continuar por indisponibilidade externa ou ausência de arquivo necessário.'),
    ]
    add_table(doc, ['Resultado', 'Critério'], criteria, [Inches(1.35), Inches(5.45)])

    doc.add_heading('Credenciais da amostra em produção', level=1)
    doc.add_paragraph(
        'O CPF pode ser digitado com ou sem pontuação. As senhas diferem por perfil para facilitar '
        'a execução dos testes e não devem ser reutilizadas em contas pessoais ou institucionais.'
    )
    credentials = [
        ('Coordenador', 'Ana Souza', '100.000.001-08', 'Coord@2026'),
        ('Coordenador', 'Bruno Lima', '100.000.002-80', 'Coord@2026'),
        ('Coordenador', 'Carla Mendes', '100.000.003-61', 'Coord@2026'),
        ('Aluno', 'Diego Alves', '100.000.004-42', 'Aluno@2026'),
        ('Aluno', 'Elisa Rocha', '100.000.005-23', 'Aluno@2026'),
        ('Aluno', 'Felipe Nunes', '100.000.006-04', 'Aluno@2026'),
        ('Gestor', 'Gabriela Costa', '100.000.007-95', 'Gestor@2026'),
        ('Gestor', 'Henrique Silva', '100.000.008-76', 'Gestor@2026'),
        ('Gestor', 'Isabela Martins', '100.000.009-57', 'Gestor@2026'),
    ]
    add_table(
        doc,
        ['Perfil', 'Nome', 'CPF', 'Senha'],
        credentials,
        [Inches(1.25), Inches(2.1), Inches(1.75), Inches(1.7)],
    )

    doc.add_heading('Caminhos do ambiente de produção', level=1)
    routes = [
        ('Entrada e autenticação', LOGIN_URL),
        ('Consulta pública', f'{PRODUCTION_URL}/projetos/consulta/'),
        ('Ajuda da plataforma', f'{PRODUCTION_URL}/projetos/ajuda/'),
        ('Painel do coordenador', f'{PRODUCTION_URL}/projetos/telaprincipal/'),
        ('Projetos do coordenador', f'{PRODUCTION_URL}/projetos/projetos/meusprojetos/'),
        ('Projetos do aluno', f'{PRODUCTION_URL}/projetos/aluno/projetos/meusprojetos/'),
        ('Painel do gestor', f'{PRODUCTION_URL}/projetos/gestor/dashboard/'),
        ('Usuários para o gestor', f'{PRODUCTION_URL}/projetos/gestor/usuarios/'),
        ('Histórico de projetos', f'{PRODUCTION_URL}/projetos/gestor/projetos/historico/'),
        ('Gestão de editais', f'{PRODUCTION_URL}/projetos/gestor/editais/'),
        ('Relatórios para o gestor', f'{PRODUCTION_URL}/projetos/gestor/relatorios/'),
        ('Administração Django', f'{PRODUCTION_URL}/admin/'),
    ]
    add_table(doc, ['Função', 'URL'], routes, [Inches(2.2), Inches(4.6)])
    doc.add_paragraph(
        'As rotas autenticadas redirecionam para a tela inicial quando a sessão não possui o perfil '
        'necessário. O painel Django usa a conta administrativa configurada no Render; as contas de '
        'gestor desta amostra devem ser usadas na própria plataforma.'
    )

    doc.add_heading('Projetos disponíveis para consulta', level=1)
    projects = [
        ('Ana Souza', 'Plataforma Aberta de Indicadores Acadêmicos', 'Agência de Fomento', 'Rascunho'),
        ('Ana Souza', 'Inclusão Digital para Comunidades do Acre', 'UFAC sem financiamento', 'Aguardando Centro'),
        ('Ana Souza', 'Laboratório de Dados para Gestão Universitária', 'UFAC com financiamento', 'Submetido'),
        ('Bruno Lima', 'Saúde Mental e Permanência Estudantil', 'Agência de Fomento', 'Aprovado com ética pendente'),
        ('Bruno Lima', 'Memória Social e Patrimônio Cultural Acreano', 'UFAC com financiamento', 'Em andamento'),
        ('Bruno Lima', 'Cartografia Participativa em Comunidades Urbanas', 'UFAC sem financiamento', 'Reprovado'),
        ('Carla Mendes', 'Manejo Sustentável de Sistemas Agroflorestais', 'UFAC sem financiamento', 'Aguardando finalização'),
        ('Carla Mendes', 'Monitoramento da Biodiversidade no Campus', 'Agência de Fomento', 'Finalizado'),
        ('Carla Mendes', 'Educação Ambiental em Escolas Públicas', 'UFAC com financiamento', 'Encerrado sem conclusão'),
    ]
    add_table(
        doc,
        ['Coordenador', 'Projeto', 'Tipo', 'Situação'],
        projects,
        [Inches(1.25), Inches(2.7), Inches(1.55), Inches(1.3)],
    )

    doc.add_heading('Dados auxiliares para os testes', level=1)
    aux_rows = [
        ('Link restrito do Centro', f'{PRODUCTION_URL}/projetos/centro/aprovacao/demo-centro-ufac-2026/'),
        ('Validade do link do Centro', 'Uso único; expira em 7 dias (168 horas) após a emissão'),
        ('Edital aberto', 'Edital DEMO nº 1 do ano corrente'),
        ('PDF nativamente digital', 'Arquivo local output/pdf/projeto_teste_preenchimento_automatico.pdf'),
        ('Imagem para anexos', 'Arquivo local static/imagens/imagem-exemplo.jpeg'),
        ('E-mail de recebimento', 'Substituir o e-mail de uma conta de amostra por uma caixa postal controlada pelo testador'),
        (
            'Centro fictício local',
            'TESTE LUNAR 20260916 - CENTRO DE ESTUDOS FICTÍCIO (SOMENTE TESTE LOCAL - NÃO USAR EM PRODUÇÃO); '
            'e-mail d.moura250304@gmail.com',
        ),
        ('CPF válido para membro manual', '529.982.247-25'),
        ('Centro sugerido', 'CCET - Centro de Ciências Exatas e Tecnológicas'),
        ('Curso sugerido', 'Bacharelado em Sistemas de Informação'),
        ('ODS sugeridos', 'ODS 4 Educação de Qualidade e ODS 9 Indústria Inovação e Infraestrutura'),
    ]
    add_table(doc, ['Item', 'Valor'], aux_rows, [Inches(2.15), Inches(4.65)])

    doc.add_heading('Dados para criar um projeto durante a validação', level=2)
    new_project_data = [
        ('Título', 'Sistema de Indicadores para Gestão Acadêmica'),
        ('Período', '01/10/2026 a 30/09/2027'),
        ('Valor de fomento', 'R$ 50.000,00'),
        ('Agência', 'CNPq'),
        ('Resumo', 'Projeto aplicado ao apoio de decisões acadêmicas por meio da organização e análise de indicadores.'),
        ('Introdução e justificativa', 'A gestão acadêmica produz dados que podem apoiar decisões, desde que tratados com segurança e transparência.'),
        ('Objetivo geral', 'Desenvolver e avaliar um protótipo de apoio à gestão acadêmica.'),
        ('Objetivos específicos', 'Mapear processos; estruturar dados; desenvolver o protótipo; avaliar resultados.'),
        ('Metodologia', 'Pesquisa aplicada, levantamento de requisitos, prototipação iterativa e avaliação com usuários.'),
        ('Resultados esperados', 'Protótipo funcional, documentação técnica e relatório de avaliação.'),
        ('Referências', 'UNIVERSIDADE FEDERAL DO ACRE. Documentos institucionais. 2026.'),
        ('Palavras-chave', 'indicadores, gestão acadêmica, inovação'),
        ('Membro manual', 'Marina Externa; CPF 529.982.247-25; Estudante Graduação; 8 h semanais; 192 h totais'),
    ]
    add_table(doc, ['Campo', 'Valor sugerido'], new_project_data, [Inches(1.75), Inches(5.05)])

    access_cases = [
        {
            'id': 'CT-ACC-001', 'title': 'Abrir a página inicial', 'profile': 'Público',
            'pre': 'Implantação do Render disponível e sessão encerrada.',
            'steps': ['Acessar a URL inicial.', 'Verificar os painéis de login, cadastro, recuperação de senha e consulta pública.'],
            'expected': ['A página carrega sem autenticação.', 'Os controles são legíveis e não há sobreposição ou rolagem horizontal indevida.'],
        },
        {
            'id': 'CT-ACC-002', 'title': 'Cadastrar usuário com máscara automática de CPF', 'profile': 'Público',
            'pre': 'Usar CPF válido ainda não cadastrado.', 'data': 'CPF 52998224725; perfil Aluno; e-mail novo e senha forte.',
            'steps': ['Abrir Criar uma conta.', 'Digitar os 11 números do CPF sem pontuação.', 'Preencher os demais campos e confirmar o cadastro.'],
            'expected': ['A máscara surge durante a digitação no formato 000.000.000-00.', 'A conta é criada ativa e o sistema orienta o login.'],
        },
        {
            'id': 'CT-ACC-003', 'title': 'Rejeitar CPF inválido', 'profile': 'Público',
            'pre': 'Tela de cadastro aberta.', 'data': 'CPF 111.111.111-11.',
            'steps': ['Preencher os demais campos válidos.', 'Informar o CPF inválido.', 'Confirmar o cadastro.'],
            'expected': ['O cadastro não é salvo.', 'Uma mensagem informa que o CPF é inválido.'],
        },
        {
            'id': 'CT-ACC-004', 'title': 'Impedir CPF ou e-mail duplicado', 'profile': 'Público',
            'pre': 'Dados de amostra carregados em produção.', 'data': 'CPF ou e-mail de Ana Souza.',
            'steps': ['Tentar criar nova conta reutilizando o CPF existente.', 'Repetir o teste reutilizando apenas o e-mail existente.'],
            'expected': ['As duas tentativas são rejeitadas sem criar usuários duplicados.', 'O campo responsável apresenta mensagem de validação.'],
        },
        {
            'id': 'CT-ACC-005', 'title': 'Autenticar os três perfis', 'profile': 'Coordenador Aluno Gestor',
            'pre': 'Dados de amostra carregados em produção.', 'data': 'Usar uma credencial de cada perfil da tabela.',
            'steps': ['Entrar como Coordenador e observar a tela principal.', 'Sair e repetir como Aluno.', 'Sair e repetir como Gestor.'],
            'expected': ['Cada credencial é aceita.', 'Menus, atalhos e painel inicial mudam conforme o perfil.'],
        },
        {
            'id': 'CT-ACC-006', 'title': 'Rejeitar senha incorreta e CPF inexistente', 'profile': 'Público',
            'pre': 'Tela de login aberta.',
            'steps': ['Tentar entrar com CPF existente e senha incorreta.', 'Tentar entrar com CPF válido não cadastrado.'],
            'expected': ['O acesso é negado nas duas tentativas.', 'A mensagem não revela detalhes sensíveis da conta.'],
        },
        {
            'id': 'CT-ACC-007', 'title': 'Validar a opção manter login', 'profile': 'Qualquer usuário',
            'pre': 'Conta ativa.',
            'steps': ['Entrar sem marcar Lembrar-me e encerrar o navegador.', 'Repetir marcando Lembrar-me.'],
            'expected': ['Sem a opção, a sessão expira ao fechar o navegador.', 'Com a opção, a sessão é configurada para até 30 dias.'],
        },
        {
            'id': 'CT-ACC-008', 'title': 'Solicitar redefinição de senha', 'profile': 'Público',
            'pre': 'E-mail transacional do Render configurado e conta de amostra apontando para uma caixa postal controlada.', 'data': 'CPF de um usuário ativo.',
            'steps': ['Abrir Esqueceu a senha.', 'Informar o CPF e enviar.', 'Abrir o link recebido e definir uma nova senha.'],
            'expected': ['A mensagem de resposta não confirma se o CPF existe.', 'O link válido permite alterar a senha e autenticar com o novo valor.'],
        },
        {
            'id': 'CT-ACC-009', 'title': 'Encerrar sessão', 'profile': 'Qualquer usuário',
            'pre': 'Usuário autenticado.',
            'steps': ['Clicar em Sair.', 'Tentar abrir diretamente uma URL autenticada.'],
            'expected': ['A sessão é encerrada.', 'A rota protegida redireciona para o acesso.'],
        },
        {
            'id': 'CT-ACC-010', 'title': 'Consultar o guia e a seção sobre o sistema', 'profile': 'Qualquer usuário autenticado',
            'pre': 'Usuário autenticado.',
            'steps': ['Abrir a tela Ajuda.', 'Percorrer as seções do guia.', 'Abrir a seção Sobre o Sistema e os Desenvolvedores.'],
            'expected': ['O guia apresenta os três tipos de projeto e os fluxos atuais.', 'Douglas Moura Araújo e Áleks Sebastian de Freitas Araújo são identificados como desenvolvedores.'],
        },
    ]
    add_cases_section(
        doc,
        'Casos de teste de acesso e cadastro',
        'Estes casos validam a entrada no sistema, a criação de contas e a recuperação de acesso.',
        access_cases,
    )

    profile_cases = [
        {
            'id': 'CT-PER-001', 'title': 'Completar o perfil do coordenador', 'profile': 'Coordenador',
            'pre': 'Entrar como Ana Souza.',
            'steps': ['Abrir Meu Perfil e Editar.', 'Preencher dados pessoais, titulação, SIAPE, Lattes e centro de lotação.', 'Salvar e reabrir o perfil.'],
            'expected': ['Os dados válidos são persistidos.', 'Centro e demais seleções exibem os valores salvos.'],
        },
        {
            'id': 'CT-PER-002', 'title': 'Restringir o regime de trabalho', 'profile': 'Coordenador',
            'pre': 'Tela de edição do perfil aberta.',
            'steps': ['Abrir Regime de Trabalho.', 'Verificar todas as opções.', 'Salvar sucessivamente 20h, 40h e DE.'],
            'expected': ['Somente 20h, 40h e DE Dedicação Exclusiva aparecem.', 'Não é possível digitar um valor livre.'],
        },
        {
            'id': 'CT-PER-003', 'title': 'Editar o perfil do aluno', 'profile': 'Aluno',
            'pre': 'Entrar como Diego Alves.',
            'steps': ['Abrir a edição de perfil.', 'Selecionar curso de graduação e vínculo de pós-graduação, se aplicável.', 'Salvar e verificar o resumo do perfil.'],
            'expected': ['A lista apresenta os cursos do catálogo UFAC.', 'Curso, matrícula e demais dados ficam disponíveis no perfil.'],
        },
        {
            'id': 'CT-PER-004', 'title': 'Validar centros e cursos vinculados', 'profile': 'Coordenador Aluno',
            'pre': 'Dados de amostra carregados em produção.',
            'steps': ['Abrir uma seleção de centro.', 'Confirmar os oito centros acadêmicos.', 'Na criação de projeto, selecionar CCET e abrir a lista de cursos.'],
            'expected': ['Não aparecem centros genéricos ou de teste.', 'Após escolher CCET, somente os quatro cursos vinculados ao centro são oferecidos.'],
        },
    ]
    add_cases_section(doc, 'Casos de teste de perfil', 'Validação dos dados pessoais e vínculos institucionais.', profile_cases)

    coordinator_cases = [
        {
            'id': 'CT-COO-001', 'title': 'Consultar painel e projetos do coordenador', 'profile': 'Coordenador',
            'pre': 'Entrar como Ana Souza.',
            'steps': ['Abrir Tela Principal.', 'Abrir Meus Projetos e alternar entre as abas de situação.'],
            'expected': ['Métricas e projetos pertencem ao coordenador autenticado.', 'O rascunho, o projeto aguardando Centro e o projeto submetido aparecem nas abas corretas.'],
        },
        {
            'id': 'CT-COO-002', 'title': 'Escolher entre rascunho e novo projeto', 'profile': 'Coordenador',
            'pre': 'Ana Souza possui um rascunho.',
            'steps': ['Clicar em Criar Novo Projeto.', 'Verificar o alerta e as ações.', 'Selecionar Iniciar um novo projeto.'],
            'expected': ['A caixa informativa mantém altura proporcional ao conteúdo.', 'O sistema informa que o novo projeto não usará os rascunhos existentes e solicita confirmação.'],
        },
        {
            'id': 'CT-COO-003', 'title': 'Criar projeto de Agência de Fomento', 'profile': 'Coordenador',
            'pre': 'Novo projeto iniciado.', 'data': 'Dados sugeridos neste documento; CNPq; R$ 50.000,00.',
            'steps': ['Selecionar Projeto Aprovado Agência de Fomento.', 'Preencher a Etapa 1.', 'Anexar comprovante PDF ou imagem e avançar.'],
            'expected': ['Agência e comprovante tornam-se obrigatórios.', 'Edital UFAC e fluxo do Centro não são exigidos.'],
        },
        {
            'id': 'CT-COO-004', 'title': 'Validar arquivo da Agência de Fomento', 'profile': 'Coordenador',
            'pre': 'Tipo Agência de Fomento selecionado.',
            'steps': ['Passar o mouse no ícone de ajuda do comprovante.', 'Tentar enviar arquivo TXT.', 'Enviar PDF, PNG ou JPG.'],
            'expected': ['O apoio informa os formatos permitidos.', 'TXT é rejeitado e PDF ou imagem é aceito.'],
        },
        {
            'id': 'CT-COO-005', 'title': 'Criar projeto UFAC sem financiamento', 'profile': 'Coordenador',
            'pre': 'Novo projeto iniciado.',
            'steps': ['Selecionar UFAC Sem Financiamento Fluxo Contínuo.', 'Escolher um centro com e-mail e preencher as etapas.', 'Submeter o projeto.'],
            'expected': ['Ata não é exigida do coordenador durante o cadastro.', 'O projeto fica Aguardando aprovação do Centro e não passa imediatamente ao gestor.'],
        },
        {
            'id': 'CT-COO-006', 'title': 'Criar projeto UFAC com financiamento', 'profile': 'Coordenador',
            'pre': 'Edital DEMO aberto.',
            'steps': ['Selecionar UFAC Com Financiamento.', 'Tentar avançar sem edital.', 'Selecionar o edital aberto e concluir a Etapa 1.'],
            'expected': ['O avanço sem edital é bloqueado.', 'Somente editais abertos podem ser escolhidos.'],
        },
        {
            'id': 'CT-COO-007', 'title': 'Declarar aprovação ética definitiva', 'profile': 'Coordenador',
            'pre': 'Na Etapa 1, marcar que o projeto envolve aspectos éticos.',
            'steps': ['Selecionar Comprovante de aprovação.', 'Anexar PDF ou imagem.', 'Concluir a etapa.'],
            'expected': ['Tipo ético, situação e arquivo correspondente são exigidos.', 'Os detalhes exibem visualmente que a aprovação definitiva foi anexada.'],
        },
        {
            'id': 'CT-COO-008', 'title': 'Declarar submissão ética e prazo de 90 dias', 'profile': 'Coordenador',
            'pre': 'Na Etapa 1, marcar que o projeto envolve aspectos éticos.',
            'steps': ['Selecionar Comprovante de submissão.', 'Anexar o comprovante e submeter.', 'Abrir os detalhes do projeto.'],
            'expected': ['O sistema registra prazo máximo de 90 dias.', 'A pendência e os dias restantes aparecem para coordenador e gestor.'],
        },
        {
            'id': 'CT-COO-009', 'title': 'Preencher manualmente a Etapa 2', 'profile': 'Coordenador',
            'pre': 'Projeto na Etapa 2.', 'data': 'Usar os textos sugeridos no início do documento.',
            'steps': ['Preencher todas as seções sem enviar documento.', 'Digitar texto suficiente para ultrapassar a altura inicial.', 'Avançar.'],
            'expected': ['O documento é opcional.', 'As caixas quebram linha, crescem com o conteúdo e preservam os valores.'],
        },
        {
            'id': 'CT-COO-010', 'title': 'Preencher automaticamente a Etapa 2', 'profile': 'Coordenador',
            'pre': 'Projeto na Etapa 2.', 'data': 'Arquivo local output/pdf/projeto_teste_preenchimento_automatico.pdf.',
            'steps': ['Enviar o PDF nativamente digital.', 'Confirmar o preenchimento automático.', 'Revisar todas as seções.'],
            'expected': ['O sistema pergunta antes de alterar os campos.', 'Resumo, introdução, objetivos, metodologia, resultados e referências são preenchidos conforme as seções identificadas.'],
        },
        {
            'id': 'CT-COO-011', 'title': 'Rejeitar documento digitalizado sem texto', 'profile': 'Coordenador',
            'pre': 'Projeto na Etapa 2.',
            'steps': ['Enviar um PDF composto apenas por imagem digitalizada.', 'Solicitar o preenchimento automático.'],
            'expected': ['O sistema não inventa conteúdo.', 'Uma mensagem orienta o uso de arquivo nativamente digital com texto selecionável.'],
        },
        {
            'id': 'CT-COO-012', 'title': 'Aplicar limite e remover arquivo na Etapa 2', 'profile': 'Coordenador',
            'pre': 'Projeto na Etapa 2 com documento anexado.',
            'steps': ['Observar o contador de caracteres.', 'Tentar exceder 4.000 caracteres.', 'Excluir o documento enviado.'],
            'expected': ['O limite de 4.000 é visível e respeitado.', 'A exclusão solicita confirmação e remove o arquivo da interface.'],
        },
        {
            'id': 'CT-COO-013', 'title': 'Adicionar usuário existente à equipe', 'profile': 'Coordenador',
            'pre': 'Projeto na Etapa 3.',
            'steps': ['Selecionar Usuário do sistema.', 'Buscar Diego por nome, curso e CPF em tentativas separadas.', 'Selecionar o resultado e informar cargas horárias.'],
            'expected': ['A busca encontra o mesmo usuário pelos três critérios.', 'Nome e CPF são preenchidos e o membro pode ser salvo sem duplicação.'],
        },
        {
            'id': 'CT-COO-014', 'title': 'Adicionar membro manual à equipe', 'profile': 'Coordenador',
            'pre': 'Projeto na Etapa 3.', 'data': 'Marina Externa; CPF 52998224725; 8 h semanais; 192 h totais.',
            'steps': ['Selecionar Preenchimento manual.', 'Digitar os dados e observar o CPF.', 'Testar todas as funções disponíveis e salvar uma delas.'],
            'expected': ['O CPF recebe máscara automaticamente.', 'As sete funções previstas estão disponíveis e o membro é salvo.'],
        },
        {
            'id': 'CT-COO-015', 'title': 'Validar a tabela da equipe em telas menores', 'profile': 'Coordenador',
            'pre': 'Etapa 3 com três ou mais linhas.',
            'steps': ['Reduzir a janela para aproximadamente 390 px.', 'Percorrer os campos e remover uma linha.'],
            'expected': ['A tabela não ultrapassa a região branca.', 'O conteúdo permanece acessível por adaptação ou rolagem interna controlada.'],
        },
        {
            'id': 'CT-COO-016', 'title': 'Selecionar vínculos com ODS', 'profile': 'Coordenador',
            'pre': 'Projeto na etapa de ODS.',
            'steps': ['Selecionar ODS 4 e ODS 9.', 'Salvar rascunho, sair e retomar.'],
            'expected': ['As seleções são persistidas.', 'Os ODS aparecem na revisão e nos detalhes do projeto.'],
        },
        {
            'id': 'CT-COO-017', 'title': 'Revisar e submeter o projeto', 'profile': 'Coordenador',
            'pre': 'Etapas obrigatórias preenchidas.',
            'steps': ['Abrir a revisão final.', 'Conferir tipo, conteúdo, equipe e ODS.', 'Submeter.'],
            'expected': ['A revisão resume os dados das etapas.', 'O status e o destino da submissão seguem o tipo escolhido.'],
        },
        {
            'id': 'CT-COO-018', 'title': 'Gerar relatório de submissão com anexos', 'profile': 'Coordenador Gestor',
            'pre': 'Projeto submetido com PDF e imagem em Anexos.',
            'steps': ['Abrir o relatório automático.', 'Conferir capa, logo, seções, equipe e ODS.', 'Verificar imagens, páginas de PDF e links dos arquivos originais.'],
            'expected': ['O PDF segue a apresentação institucional e contém os dados do projeto.', 'Anexos válidos são incorporados e permanecem acessíveis por link.'],
        },
        {
            'id': 'CT-COO-019', 'title': 'Salvar retomar e excluir rascunho', 'profile': 'Coordenador',
            'pre': 'Novo projeto parcialmente preenchido.',
            'steps': ['Salvar como rascunho e sair.', 'Retomar pela aba Rascunhos.', 'Excluir o rascunho após confirmar a ação.'],
            'expected': ['Os dados permanecem ao retomar.', 'Somente o rascunho selecionado é removido.'],
        },
        {
            'id': 'CT-COO-020', 'title': 'Iniciar projeto aprovado e validar bloqueio ético', 'profile': 'Coordenador',
            'pre': 'Usar Saúde Mental e Permanência Estudantil.',
            'steps': ['Tentar iniciar com aprovação ética pendente.', 'Anexar a aprovação definitiva.', 'Tentar iniciar novamente.'],
            'expected': ['A primeira tentativa é bloqueada com orientação.', 'Após o comprovante aprovado, o início pode ser confirmado e a equipe é notificada.'],
        },
        {
            'id': 'CT-COO-021', 'title': 'Enviar relatórios e gerar certificados', 'profile': 'Coordenador',
            'pre': 'Usar projeto Em andamento ou Aguardando finalização.',
            'steps': ['Enviar relatório parcial com evidências.', 'Enviar relatório final.', 'Abrir Certificados e gerar para coordenador e membros.'],
            'expected': ['Relatórios e evidências ficam associados ao projeto.', 'O relatório final conduz ao fluxo de finalização e os certificados exibem os dados corretos.'],
        },
    ]
    add_cases_section(
        doc,
        'Casos de teste do coordenador',
        'Os testes percorrem as cinco etapas atuais do cadastro e as ações posteriores à submissão.',
        coordinator_cases,
    )

    center_cases = [
        {
            'id': 'CT-CEN-001', 'title': 'Abrir o link restrito do Centro', 'profile': 'Responsável do Centro',
            'pre': 'Usar um projeto temporário local ou outro token criado para o teste. Não consumir o link público de demonstração em verificações comuns.',
            'steps': ['Abrir o link em janela anônima, sem autenticação.', 'Conferir o resumo do projeto.'],
            'expected': ['O acesso exibe apenas as informações necessárias.', 'Menus e funções internas da plataforma não ficam disponíveis.'],
        },
        {
            'id': 'CT-CEN-002', 'title': 'Visualizar o documento do projeto', 'profile': 'Responsável do Centro',
            'pre': 'Link do Centro ativo.',
            'steps': ['Clicar para visualizar o projeto.', 'Voltar ao formulário de aprovação.'],
            'expected': ['O documento pode ser consultado pelo token.', 'Nenhum outro projeto pode ser acessado pelo mesmo link.'],
        },
        {
            'id': 'CT-CEN-003', 'title': 'Validar campos obrigatórios da aprovação', 'profile': 'Responsável do Centro',
            'pre': 'Link do Centro ativo.',
            'steps': ['Enviar sem nome, cargo, confirmação ou ata.', 'Tentar anexar formato não permitido.'],
            'expected': ['Campos ausentes são indicados.', 'Somente PDF, PNG, JPG ou JPEG são aceitos para a ata.'],
        },
        {
            'id': 'CT-CEN-004', 'title': 'Registrar a ata e encaminhar à PROPEG', 'profile': 'Responsável do Centro',
            'pre': 'Link ativo e arquivo de teste disponível.',
            'steps': [
                'Informar nome e cargo.',
                'Confirmar a aprovação, anexar a ata e enviar.',
                'Abrir os detalhes do projeto com o coordenador ou gestor e consultar a ata em Anexos.',
            ],
            'expected': [
                'A ata fica vinculada ao projeto e pode ser aberta em Anexos.',
                'O projeto muda de Aguardando aprovação do Centro para Submetido, segue à gestão e o token é marcado como utilizado.',
                'O coordenador recebe a notificação prevista.',
            ],
        },
        {
            'id': 'CT-CEN-005', 'title': 'Impedir reutilização ou acesso expirado', 'profile': 'Responsável do Centro',
            'pre': 'Token utilizado no caso anterior ou token expirado.',
            'steps': ['Abrir novamente o mesmo link.', 'Tentar enviar outra ata.'],
            'expected': ['O projeto não é reaberto para alteração.', 'A página informa que o link foi utilizado, expirou ou foi invalidado.'],
        },
        {
            'id': 'CT-CEN-006', 'title': 'Enviar ata quando o navegador informa origem opaca', 'profile': 'Responsável do Centro',
            'pre': 'Token local válido e cliente configurado para enviar o cabeçalho HTTP Origin: null, sem Referer.',
            'steps': [
                'Abrir o link público sem autenticação.',
                'Preencher nome, cargo e confirmação, anexar uma ata válida e enviar com Origin: null.',
                'Reabrir o mesmo link e consultar o projeto no perfil autorizado.',
            ],
            'expected': [
                'O envio não é recusado pela verificação CSRF.',
                'A ata é salva, o token é consumido uma única vez e o projeto passa para Submetido.',
                'O segundo acesso informa que o link está indisponível.',
            ],
        },
    ]
    add_cases_section(doc, 'Casos de teste do Centro de Estudos', 'O link é temporário, restrito e de uso único.', center_cases)

    manager_cases = [
        {
            'id': 'CT-GES-001', 'title': 'Consultar o painel do gestor', 'profile': 'Gestor',
            'pre': 'Entrar como Gabriela Costa.',
            'steps': ['Abrir o Painel do Gestor.', 'Conferir métricas e projetos que aguardam ação.'],
            'expected': ['Os totais refletem os dados de amostra em produção.', 'A navegação permite acessar usuários, projetos, relatórios e editais.'],
        },
        {
            'id': 'CT-GES-002', 'title': 'Pesquisar e consultar usuários', 'profile': 'Gestor',
            'pre': 'Dados de amostra carregados em produção.',
            'steps': ['Abrir Gerenciar Usuários.', 'Pesquisar por nome, CPF e perfil.', 'Abrir os detalhes de um resultado.'],
            'expected': ['Os filtros retornam usuários compatíveis.', 'Detalhes pessoais e institucionais aparecem conforme o cadastro.'],
        },
        {
            'id': 'CT-GES-003', 'title': 'Ativar inativar e editar usuário', 'profile': 'Gestor',
            'pre': 'Usuário de amostra selecionado.',
            'steps': ['Usar a ação de inativar e confirmar.', 'Tentar login com a conta.', 'Reativar e editar os campos permitidos.'],
            'expected': ['Conta inativa não autentica.', 'A reativação e a edição ficam registradas e o usuário é notificado.'],
        },
        {
            'id': 'CT-GES-004', 'title': 'Excluir usuário sem vínculos e proteger usuário vinculado', 'profile': 'Gestor',
            'pre': 'Criar uma conta temporária sem projetos; manter os usuários de amostra vinculados.',
            'steps': ['Excluir a conta temporária após confirmar.', 'Tentar excluir coordenador ou aluno com projeto.'],
            'expected': ['A conta sem vínculos é excluída.', 'Vínculos protegidos impedem exclusão inconsistente e geram orientação.'],
        },
        {
            'id': 'CT-GES-005', 'title': 'Analisar detalhes de projeto submetido', 'profile': 'Gestor',
            'pre': 'Usar Laboratório de Dados para Gestão Universitária.',
            'steps': ['Abrir os detalhes.', 'Conferir tipo, edital, conteúdo, equipe, ODS, anexos e auditoria.'],
            'expected': ['Todas as informações necessárias à decisão estão disponíveis.', 'O edital e o total de submissões podem ser consultados.'],
        },
        {
            'id': 'CT-GES-006', 'title': 'Aprovar projeto', 'profile': 'Gestor',
            'pre': 'Projeto submetido sem pendências.',
            'steps': ['Clicar em Aprovar.', 'Ler o apoio de confirmação e confirmar.'],
            'expected': ['O status muda para Aprovado.', 'Coordenador recebe notificação interna e e-mail.'],
        },
        {
            'id': 'CT-GES-007', 'title': 'Reprovar projeto e permitir ressubmissão', 'profile': 'Gestor Coordenador',
            'pre': 'Projeto submetido.',
            'steps': ['Reprovar e informar a justificativa, quando solicitada.', 'Entrar como coordenador e abrir o projeto.', 'Editar e ressubmeter.'],
            'expected': ['O status muda para Reprovado e o coordenador é notificado.', 'O coordenador consegue corrigir e reenviar o projeto.'],
        },
        {
            'id': 'CT-GES-008', 'title': 'Bloquear aprovação com ética pendente', 'profile': 'Gestor',
            'pre': 'Usar Saúde Mental e Permanência Estudantil.',
            'steps': ['Observar o alerta visual de pendência.', 'Tentar aprovar ou iniciar o fluxo sem o comprovante definitivo.'],
            'expected': ['A pendência é destacada nos detalhes.', 'A ação é bloqueada até a aprovação ética ser anexada.'],
        },
        {
            'id': 'CT-GES-009', 'title': 'Corrigir dados do projeto e validar auditoria', 'profile': 'Gestor',
            'pre': 'Projeto existente selecionado.',
            'steps': ['Abrir Editar Projeto.', 'Alterar um campo e substituir ou remover um documento.', 'Salvar e voltar aos detalhes.'],
            'expected': ['A correção é persistida.', 'Data, hora e nome do gestor da última alteração aparecem visualmente.'],
        },
        {
            'id': 'CT-GES-010', 'title': 'Cadastrar projeto existente', 'profile': 'Gestor',
            'pre': 'Dados do projeto e coordenador disponíveis.',
            'steps': ['Abrir Cadastrar Projeto Existente.', 'Selecionar coordenador e percorrer todas as etapas.', 'Informar a situação atual e concluir.'],
            'expected': ['O registro histórico é criado com todos os dados.', 'O projeto aparece no histórico com indicação de importação legada.'],
        },
        {
            'id': 'CT-GES-011', 'title': 'Filtrar e ordenar o histórico de projetos', 'profile': 'Gestor',
            'pre': 'Dados de amostra carregados em produção.',
            'steps': ['Combinar busca textual, status, centro, curso, datas, tipo e múltiplos ODS.', 'Testar o modo Todas as ODS (AND) e depois Qualquer ODS (OR).', 'Alterar ordenação e itens por página.'],
            'expected': ['Filtros são combinados sem perder parâmetros na paginação; AND exige todas as ODS e OR aceita qualquer uma.', 'A tabela e os totais correspondem aos critérios escolhidos.'],
        },
        {
            'id': 'CT-GES-012', 'title': 'Exportar o histórico em PDF', 'profile': 'Gestor',
            'pre': 'Aplicar um conjunto de filtros no histórico.',
            'steps': ['Acionar a exportação em PDF.', 'Comparar resultados com a tela.'],
            'expected': ['O PDF é gerado sem erro.', 'Projetos e filtros relevantes correspondem à consulta exibida.'],
        },
        {
            'id': 'CT-GES-013', 'title': 'Criar editar publicar e fechar edital', 'profile': 'Gestor',
            'pre': 'Dados de edital e PDF opcional disponíveis.',
            'steps': ['Criar edital como rascunho.', 'Editar e publicar.', 'Fechar e reabrir usando a ação confirmada.'],
            'expected': ['Status e período de submissão são respeitados.', 'Somente editais abertos aparecem para novos projetos UFAC com financiamento.'],
        },
        {
            'id': 'CT-GES-014', 'title': 'Gerenciar adendos e documentos de edital', 'profile': 'Gestor',
            'pre': 'Edital existente.',
            'steps': ['Adicionar um adendo com descrição e arquivo.', 'Abrir o edital como coordenador.', 'Excluir o adendo após confirmação.'],
            'expected': ['O adendo fica disponível durante sua existência.', 'A exclusão remove apenas o adendo selecionado.'],
        },
        {
            'id': 'CT-GES-015', 'title': 'Consultar quantidade de projetos por edital', 'profile': 'Gestor',
            'pre': 'Edital DEMO nº 1 possui projetos vinculados.',
            'steps': ['Abrir a lista e os detalhes do edital.', 'Verificar total de projetos submetidos e acessar a relação.'],
            'expected': ['O total é exibido de forma clara.', 'A contagem considera projetos UFAC com financiamento vinculados ao edital.'],
        },
        {
            'id': 'CT-GES-016', 'title': 'Acompanhar relatórios de projetos', 'profile': 'Gestor',
            'pre': 'Projeto com relatório parcial ou final.',
            'steps': ['Abrir Relatórios.', 'Filtrar e abrir um relatório com evidências.'],
            'expected': ['O gestor visualiza responsável, projeto, tipo, data e arquivos.', 'Evidências permanecem acessíveis.'],
        },
        {
            'id': 'CT-GES-017', 'title': 'Distinguir Finalizado de Encerrado', 'profile': 'Gestor',
            'pre': 'Usar os dois projetos finais da base.',
            'steps': ['Abrir Monitoramento da Biodiversidade no Campus.', 'Abrir Educação Ambiental em Escolas Públicas.', 'Comparar status, datas e motivo.'],
            'expected': ['Finalizado representa conclusão com êxito.', 'Encerrado sem conclusão exibe o motivo e o gestor responsável.'],
        },
    ]
    add_cases_section(doc, 'Casos de teste do gestor', 'Validação das rotinas administrativas, de decisão e de auditoria.', manager_cases)

    student_cases = [
        {
            'id': 'CT-ALU-001', 'title': 'Consultar painel do aluno', 'profile': 'Aluno',
            'pre': 'Entrar como Diego Alves.',
            'steps': ['Abrir Tela Principal e Meus Projetos.', 'Comparar a lista com as participações da base.'],
            'expected': ['Somente projetos em que Diego integra a equipe aparecem.', 'Status e coordenador são exibidos corretamente.'],
        },
        {
            'id': 'CT-ALU-002', 'title': 'Visualizar detalhes permitidos', 'profile': 'Aluno',
            'pre': 'Aluno vinculado a projeto.',
            'steps': ['Abrir um projeto da lista.', 'Conferir resumo, objetivo, equipe e ODS.'],
            'expected': ['A visualização simplificada contém informações pertinentes.', 'Dados administrativos restritos e ações de edição não aparecem.'],
        },
        {
            'id': 'CT-ALU-003', 'title': 'Bloquear criação e edição de projetos', 'profile': 'Aluno',
            'pre': 'Aluno autenticado.',
            'steps': ['Tentar acessar diretamente a URL de criação.', 'Tentar abrir a edição de um projeto conhecido.'],
            'expected': ['O sistema nega ou redireciona o acesso.', 'Nenhum dado do projeto é alterado.'],
        },
        {
            'id': 'CT-ALU-004', 'title': 'Consultar perfil do coordenador', 'profile': 'Aluno',
            'pre': 'Projeto da equipe aberto.',
            'steps': ['Abrir o perfil público interno do coordenador.'],
            'expected': ['Informações profissionais permitidas são exibidas.', 'Dados privados não são expostos indevidamente.'],
        },
        {
            'id': 'CT-ALU-005', 'title': 'Acessar certificado após conclusão', 'profile': 'Aluno',
            'pre': 'Aluno participante de projeto finalizado.',
            'steps': ['Abrir o projeto finalizado.', 'Acessar ou solicitar o certificado disponível.'],
            'expected': ['O certificado identifica projeto, participante, função e cargas horárias.', 'Aluno não consegue gerar certificado de pessoa não autorizada.'],
        },
    ]
    add_cases_section(
        doc,
        'Casos de teste do aluno',
        'Validação do acompanhamento e das restrições do perfil discente.',
        student_cases,
    )

    public_cases = [
        {
            'id': 'CT-PUB-001', 'title': 'Abrir consulta pública sem conta', 'profile': 'Público',
            'pre': 'Sessão encerrada.',
            'steps': ['Acessar Consulta Pública.', 'Navegar pela lista de projetos.'],
            'expected': ['Não é solicitado login.', 'Somente situações publicáveis aparecem.'],
        },
        {
            'id': 'CT-PUB-002', 'title': 'Pesquisar e filtrar consulta pública', 'profile': 'Público',
            'pre': 'Consulta pública aberta.',
            'steps': ['Pesquisar por título ou coordenador.', 'Selecionar ODS 1 e ODS 10.', 'Testar Todas as ODS (AND) e depois Qualquer ODS (OR).'],
            'expected': ['No modo AND aparecem somente projetos vinculados às duas ODS; no modo OR aparece qualquer projeto vinculado a pelo menos uma.', 'Filtros podem ser limpos sem erro.'],
        },
        {
            'id': 'CT-PUB-003', 'title': 'Visualizar detalhes públicos', 'profile': 'Público',
            'pre': 'Projeto publicável selecionado.',
            'steps': ['Abrir os detalhes.', 'Revisar campos e links disponíveis.'],
            'expected': ['Resumo, objetivos, coordenador, centro e ODS são exibidos.', 'CPF, documentos restritos e ações administrativas não aparecem.'],
        },
        {
            'id': 'CT-PUB-004', 'title': 'Exportar consulta pública em PDF', 'profile': 'Público',
            'pre': 'Consulta com filtros aplicados.',
            'steps': ['Gerar o PDF.', 'Comparar o conteúdo com a tela.'],
            'expected': ['O arquivo é gerado e legível.', 'O resultado respeita os filtros aplicados.'],
        },
    ]
    add_cases_section(
        doc,
        'Casos de teste da consulta pública',
        'Validação do acesso sem autenticação e da proteção de dados.',
        public_cases,
    )

    notification_cases = [
        {
            'id': 'CT-NOT-001', 'title': 'Visualizar e excluir notificações', 'profile': 'Qualquer usuário autenticado',
            'pre': 'Usuário possui a notificação inicial da amostra.',
            'steps': ['Abrir o sino e a central de notificações.', 'Pesquisar por texto ou data.', 'Excluir uma notificação e depois excluir múltiplas selecionadas.'],
            'expected': ['Contador e estado de leitura são atualizados.', 'Somente notificações do usuário autenticado são alteradas.'],
        },
        {
            'id': 'CT-NOT-002', 'title': 'Emitir alertas do prazo ético', 'profile': 'Coordenador Gestor',
            'pre': 'Projeto com submissão ética e prazo configurado.',
            'steps': ['Executar python manage.py notificar_prazos_etica em um marco aplicável.', 'Consultar notificações e caixa de e-mail de teste.', 'Executar novamente no mesmo marco.'],
            'expected': ['Alerta interno e e-mail são enviados em 60, 30, 15, 7, 3 e 1 dia.', 'O mesmo marco não gera duplicidade.'],
        },
    ]
    add_cases_section(
        doc,
        'Casos de teste de notificações',
        'Validação da central interna e dos avisos por prazo.',
        notification_cases,
    )

    security_cases = [
        {
            'id': 'CT-SEG-001', 'title': 'Validar autorização por perfil em URL direta', 'profile': 'Coordenador Aluno Gestor',
            'pre': 'Uma sessão ativa por vez.',
            'steps': ['Com aluno, abrir rota de gestor.', 'Com coordenador, abrir gestão de usuários.', 'Com gestor, confirmar acesso às rotas administrativas.'],
            'expected': ['Perfis sem permissão recebem resposta segura ou redirecionamento.', 'Nenhuma informação restrita é processada antes da autorização.'],
        },
        {
            'id': 'CT-SEG-002', 'title': 'Impedir acesso a projeto de outro coordenador', 'profile': 'Coordenador',
            'pre': 'Entrar como Ana e conhecer o ID de projeto de Bruno.',
            'steps': ['Tentar abrir a edição do projeto de Bruno pela URL.', 'Tentar alterar uma etapa por requisição direta.'],
            'expected': ['A edição é negada.', 'O projeto permanece sem alterações.'],
        },
        {
            'id': 'CT-SEG-003', 'title': 'Validar formatos de upload', 'profile': 'Coordenador Gestor Centro',
            'pre': 'Telas de upload disponíveis.',
            'steps': ['Tentar enviar executável ou TXT onde só PDF ou imagem é permitido.', 'Enviar PDF ou imagem válida.'],
            'expected': ['Extensões indevidas são rejeitadas no servidor, mesmo se o filtro do navegador for contornado.', 'Formatos permitidos são aceitos.'],
        },
        {
            'id': 'CT-SEG-004', 'title': 'Confirmar ações destrutivas ou críticas', 'profile': 'Coordenador Gestor',
            'pre': 'Rascunho, anexo, usuário e edital de teste disponíveis.',
            'steps': ['Acionar excluir, aprovar, reprovar, fechar ou encerrar.', 'Passar o mouse no apoio de confirmação e cancelar.', 'Repetir e confirmar.'],
            'expected': ['O apoio informa a consequência da ação.', 'Cancelar não altera dados e confirmar executa somente o alvo selecionado.'],
        },
        {
            'id': 'CT-SEG-005', 'title': 'Validar token do Centro contra alteração', 'profile': 'Público',
            'pre': 'Link de amostra do Centro disponível.',
            'steps': ['Alterar um ou mais caracteres do token na URL.', 'Tentar consultar e enviar dados.'],
            'expected': ['O token alterado não localiza a solicitação.', 'Nenhum dado do projeto é exposto.'],
        },
        {
            'id': 'CT-SEG-006', 'title': 'Validar responsividade geral', 'profile': 'Todos',
            'pre': 'Navegador com ferramentas responsivas.',
            'steps': ['Testar login, ajuda, perfis, etapas, tabelas e detalhes em 390 px, 768 px e largura desktop.', 'Verificar navegação por teclado.'],
            'expected': ['Não há conteúdo essencial cortado ou fora da região principal.', 'Campos, botões, modais e links permanecem utilizáveis.'],
        },
    ]
    add_cases_section(
        doc,
        'Casos de teste de segurança e interface',
        'Cenários negativos e de controle de acesso.',
        security_cases,
    )

    e2e_heading = doc.add_heading('Roteiros completos de ponta a ponta', level=1)
    e2e_heading.paragraph_format.page_break_before = True
    doc.add_paragraph(
        'Os roteiros seguintes conectam os casos isolados e devem ser executados em uma base recém-restaurada '
        'ou com projetos novos, para que a mudança de status não interfira nos demais testes.'
    )
    flows = [
        (
            'Fluxo A Agência de Fomento',
            [
                'Coordenador cria o projeto, escolhe Agência de Fomento e anexa a aprovação externa.',
                'Coordenador preenche conteúdo, equipe e ODS, revisa e submete.',
                'Gestor confere documentos, corrige um campo para validar auditoria e aprova.',
                'Coordenador inicia o projeto, envia relatório parcial e depois relatório final.',
                'Gestor confirma a conclusão; o status final deve ser Finalizado.',
            ],
        ),
        (
            'Fluxo B UFAC sem financiamento',
            [
                'Coordenador seleciona UFAC Sem Financiamento e um centro com e-mail.',
                'Após a submissão, o projeto fica Aguardando aprovação do Centro.',
                'Responsável abre o link restrito, consulta o projeto e anexa a ata aprovada.',
                'O token torna-se inutilizável e o projeto segue para a gestão.',
                'Gestor aprova; coordenador inicia, envia relatórios e conclui.',
            ],
        ),
        (
            'Fluxo C UFAC com financiamento',
            [
                'Gestor publica edital com período de submissão vigente.',
                'Coordenador seleciona UFAC Com Financiamento e vincula o edital aberto.',
                'Após revisão e submissão, o projeto aparece na contagem do edital.',
                'Gestor aprova ou reprova; em reprovação, coordenador corrige e ressubmete.',
                'Após aprovação e execução, relatório final conduz ao status Finalizado.',
            ],
        ),
        (
            'Fluxo D Pendência ética',
            [
                'Coordenador declara aspectos éticos e anexa apenas o comprovante de submissão.',
                'Sistema registra prazo de 90 dias e apresenta a pendência nos detalhes.',
                'Comando de alertas envia avisos nos marcos previstos sem duplicação.',
                'Gestor não consegue aprovar e coordenador não consegue iniciar enquanto faltar a aprovação.',
                'Após anexar o comprovante definitivo, o bloqueio deixa de existir.',
            ],
        ),
    ]
    for title_text, steps in flows:
        doc.add_heading(title_text, level=2)
        for number, step in enumerate(steps, start=1):
            add_numbered(doc, number, step)

    automated_heading = doc.add_heading('Verificação automatizada e manutenção da amostra', level=1)
    automated_heading.paragraph_format.page_break_before = True
    automated_intro = doc.add_paragraph(
        'Os comandos abaixo são destinados ao responsável técnico e devem ser executados no ambiente '
        'do Render ou em uma implantação equivalente. O modo aditivo preserva registros existentes; '
        'a limpeza remove somente os CPFs, projetos e editais identificados como demonstração.'
    )
    automated_heading.paragraph_format.keep_with_next = True
    automated_intro.paragraph_format.keep_with_next = True
    automated = [
        ('Verificação do projeto', 'python manage.py check', 'Nenhum problema identificado.'),
        ('Migrations pendentes', 'python manage.py makemigrations --check --dry-run', 'Nenhuma alteração detectada.'),
        ('Testes automatizados', 'python manage.py test', 'Todos os testes concluídos com OK.'),
        ('Alertas éticos', 'python manage.py notificar_prazos_etica', 'Alertas aplicáveis processados sem duplicidade.'),
        ('Adicionar ou restaurar amostra', 'python manage.py povoar_demonstracao --adicionar', '9 usuários e 9 projetos de amostra, sem excluir outros dados.'),
        ('Remover amostra temporária', 'python manage.py limpar_demonstracao --confirmar', 'Somente usuários, projetos e editais DEMO são removidos.'),
    ]
    add_table(doc, ['Objetivo', 'Comando', 'Resultado esperado'], automated, [Inches(1.5), Inches(3.2), Inches(2.1)])

    doc.add_heading('Registro de execução', level=1)
    doc.add_paragraph(
        'Use uma linha por caso executado. Quando houver falha, anexe captura de tela, informe a URL, '
        'o perfil utilizado, os dados inseridos e o horário aproximado.'
    )
    records = [('', '', '', '', '') for _ in range(12)]
    add_table(
        doc,
        ['Caso', 'Data', 'Resultado', 'Evidência', 'Observação'],
        records,
        [Inches(1.0), Inches(0.9), Inches(1.0), Inches(1.6), Inches(2.3)],
    )

    update_heading = doc.add_heading('Atualização da execução de 16 de setembro de 2026', level=1)
    update_heading.paragraph_format.page_break_before = True
    doc.add_paragraph(
        'Esta seção consolida a validação específica do fluxo de aprovação pelo Centro de Estudos, '
        'a correção aplicada durante a execução e os limites da confirmação em produção. Os resultados '
        'não autorizam declarar a plataforma livre de bugs; eles descrevem a cobertura efetivamente executada.'
    )

    doc.add_heading('Resultado do fluxo do Centro de Estudos', level=2)
    execution_rows = [
        (
            'Envio local com token válido',
            'Corrigido',
            'Ata salva; solicitação marcada como utilizada; projeto 106 passou de Aguardando aprovação do Centro para Submetido.',
        ),
        (
            'Consulta posterior da ata',
            'Passou',
            'A ata apareceu em Anexos nos detalhes do projeto e o PDF pôde ser aberto pelo perfil autorizado.',
        ),
        (
            'Reutilização do token local',
            'Passou',
            'O token consumido deixou de permitir novo envio.',
        ),
        (
            'Produção com Origin: null e token inválido',
            'Passou',
            'A requisição alcançou a regra do link e retornou 410 Link indisponível, em vez de 403 CSRF.',
        ),
        (
            'Produção com token válido e ata real',
            'Bloqueado',
            'Não executado para preservar o link de demonstração de uso único e evitar alteração de dados de produção.',
        ),
        (
            'Suíte automatizada completa',
            'Passou',
            '49 testes concluídos com sucesso; check sem problemas; nenhuma migration pendente.',
        ),
    ]
    add_table(
        doc,
        ['Verificação', 'Resultado', 'Evidência'],
        execution_rows,
        [Inches(2.0), Inches(1.05), Inches(3.75)],
    )

    doc.add_heading('BUG-001 — envio da ata bloqueado por CSRF em origem isolada', level=2)
    bug_rows = [
        ('Classificação', 'Alta'),
        ('Perfil', 'Responsável do Centro de Estudos, sem autenticação'),
        ('Página', '/projetos/centro/aprovacao/<token>/'),
        ('Pré-condição', 'Solicitação ativa para projeto UFAC sem financiamento e ata válida disponível.'),
        (
            'Passos para reproduzir',
            'Abrir o link em contexto que envie Origin: null e não envie Referer; preencher nome, cargo e confirmação; anexar a ata; enviar.',
        ),
        (
            'Resultado observado',
            'Resposta 403: “Origin checking failed - null does not match any trusted origins.” Os dados não eram processados.',
        ),
        (
            'Resultado esperado',
            'A ata deve ser validada e vinculada; o token deve ser consumido; o projeto deve seguir para Submetido.',
        ),
        (
            'Causa',
            'O middleware CSRF interrompia o POST antes da view quando o navegador usava origem opaca. Esse endpoint público não usa sessão como credencial; a autorização é feita por token aleatório, temporário e de uso único.',
        ),
        (
            'Correção',
            'Isenção de CSRF limitada à view aprovacao_centro. Permaneceram as validações de token, expiração, uso único, transação atômica, confirmação e formato do arquivo. As demais rotas continuam protegidas por CSRF.',
        ),
        (
            'Teste de regressão',
            'test_link_publico_aceita_ata_quando_navegador_envia_origin_null: falhava com 403 antes da correção e passou com ata salva, token utilizado e projeto Submetido.',
        ),
        (
            'Publicação',
            'Commit 00e815e — Corrige envio de ata em origem isolada; versão enviada à branch main e confirmada no Render.',
        ),
    ]
    add_table(doc, ['Campo', 'Registro'], bug_rows, [Inches(1.55), Inches(5.25)])

    doc.add_heading('Controles de segurança confirmados no link', level=2)
    add_bullet(doc, 'O valor bruto do token aparece somente no link enviado; o banco armazena seu resumo criptográfico SHA-256.')
    add_bullet(doc, 'O token expira em 168 horas, é de uso único e solicitações anteriores podem ser invalidadas.')
    add_bullet(doc, 'O processamento usa transação e bloqueio do registro para impedir consumo concorrente.')
    add_bullet(doc, 'Nome, cargo, confirmação e arquivo permitido são validados no servidor.')
    add_bullet(doc, 'A resposta do link usa política de referência restrita para reduzir exposição do token em navegação externa.')

    doc.add_heading('RISCO-001 — acesso direto aos arquivos armazenados', level=2)
    add_label_paragraph(doc, 'Classificação proposta', 'Alta, caso atas e anexos sejam documentos restritos.')
    add_label_paragraph(
        doc,
        'Constatação',
        'A configuração atual do armazenamento de mídia no Cloudinary permite acesso anônimo ao arquivo por quem possuir a URL direta.',
    )
    add_label_paragraph(
        doc,
        'Impacto possível',
        'O controle de perfil na página de detalhes não revoga uma URL de mídia já conhecida ou compartilhada.',
    )
    add_label_paragraph(
        doc,
        'Ação recomendada',
        'Definir a classificação institucional das atas. Se forem restritas, usar download autenticado, URLs assinadas com validade curta ou armazenamento privado e adicionar testes de autorização.',
    )
    add_label_paragraph(
        doc,
        'Situação',
        'Pendente de decisão funcional e de uma validação específica de autorização de documentos em produção.',
    )

    doc.add_heading('Dados temporários criados em ambiente local', level=2)
    temporary_rows = [
        (
            'Centro 24',
            'TESTE LUNAR 20260916 - CENTRO DE ESTUDOS FICTÍCIO (SOMENTE TESTE LOCAL - NÃO USAR EM PRODUÇÃO)',
        ),
        ('Usuário 36', 'teste-lunar-coordenador; conta sintética usada apenas no fluxo local.'),
        ('Projeto 106', 'TESTE LUNAR 20260916 - VALIDAÇÃO DO LINK DO CENTRO - TESTE LOCAL; situação final Submetido.'),
        ('Anexo 58', 'Ata fictícia vinculada ao projeto 106.'),
        ('Arquivo de exemplo', 'output/pdf/TESTE_LUNAR_ATA_FICTICIA_20260916.pdf'),
        ('Arquivo enviado', 'media/anexos/2026/09/TESTE_LUNAR_ATA_FICTICIA_20260916.pdf'),
    ]
    add_table(doc, ['Registro', 'Identificação'], temporary_rows, [Inches(1.35), Inches(5.45)])
    doc.add_paragraph(
        'Limpeza: remover somente o projeto 106 e seus anexos, notificações e solicitação de aprovação; '
        'depois remover o usuário 36 e o centro 24, confirmando antes que nenhum vínculo externo tenha sido criado. '
        'Não remover usuários, projetos ou o link de demonstração existentes antes desta execução.'
    )

    doc.add_heading('Conclusão desta execução', level=2)
    doc.add_paragraph(
        'O fluxo do Centro de Estudos foi confirmado de ponta a ponta no ambiente local após a correção. '
        'Em produção, foi confirmado que a implantação aceita o contexto Origin: null e alcança a validação '
        'do token. O envio com um token de produção válido permanece sem execução para preservar o link de '
        'demonstração. Portanto, a funcionalidade está operacional dentro da cobertura descrita, mas a plataforma '
        'não deve ser considerada livre de bugs conhecidos enquanto o RISCO-001 e os demais casos ainda não '
        'executados estiverem pendentes.'
    )

    doc.add_heading('Critério de encerramento da validação', level=1)
    doc.add_paragraph(
        'A validação pode ser encerrada quando todos os casos de prioridade alta estiverem aprovados, '
        'não houver falhas de controle de acesso ou perda de dados e as divergências restantes estiverem '
        'registradas com responsável e decisão de correção. Os três fluxos de projeto e o fluxo ético '
        'devem ser executados integralmente ao menos uma vez. O envio em produção com token válido pode ser '
        'executado somente com um projeto temporário autorizado, sem consumir o link de demonstração.'
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == '__main__':
    build_document()
